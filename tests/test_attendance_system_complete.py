# tests/test_attendance_system_complete.py

from __future__ import annotations
import os
from datetime import date, datetime
from pathlib import Path
import pytest
from PIL import Image, ImageDraw

from core.database.engine import DatabaseEngine
from modules.attendance.constants import (
    AttendanceHealth,
    AttendanceMode,
    MarkingMethod,
    AttendanceStatus,
)
from modules.attendance.dto import (
    AttendanceRecordDTO,
    AttendanceSessionDTO,
    BatchAttendanceMatrixDTO,
    BatchStudentDTO,
    OCRScanResultDTO,
    OfflineFormatConfigDTO,
    StudentAttendanceHealthDTO,
)
from modules.attendance.ocr_engine import AttendanceOCREngine
from modules.attendance.service import AttendanceService
from modules.attendance.controller import AttendanceController
from modules.attendance.views.attendance_home import AttendanceHome
from modules.attendance.views.offline_format_dialog import OfflineFormatDialog
from modules.attendance.views.offline_upload_modal import OfflineUploadModal
from ui.layout.navigation import NavigationMenu
from ui.router import AppRouter


@pytest.fixture
def setup_attendance_fixture():
    """
    Sets up a clean course, batch, and enrolled students in the isolated test database.
    """
    from core.security.context import SecurityContext
    from core.security.roles import Role
    from modules.admission.service import AdmissionService
    from modules.admission.dto import AdmissionCreateDTO, AdmissionStatus
    from modules.course.service import CourseService
    from modules.course.dto import CourseCreateDTO
    from modules.batch.controller import BatchController

    SecurityContext.set_current_user(user_id=1, username="admin", role=Role.ADMINISTRATOR.value)

    course_service = CourseService()
    batch_ctrl = BatchController()
    adm_service = AdmissionService()

    code = f"AT{int(datetime.now().timestamp() * 1000) % 100000}"
    course_id = course_service.create_course(
        CourseCreateDTO(name="Full Stack Python", code=code, base_fee=25000.0, duration="6 Months")
    )

    batch_id = batch_ctrl.create_batch({
        "course_id": course_id,
        "batch_name": f"Morning Batch {code}",
        "timing": "09:00 AM - 11:00 AM",
        "max_capacity": 30,
        "status": "OPEN",
    })

    student_ids = []
    names = ["Aarav Mehta", "Neha Patil", "Vikram Rathore"]
    for idx, name in enumerate(names, start=1):
        fn, ln = name.split()
        adm_dto = AdmissionCreateDTO(
            first_name=fn,
            last_name=ln,
            mobile_number=f"981100{idx:04d}",
            secondary_mobile=f"982200{idx:04d}",
            village="Nashik",
            course_id=course_id,
            batch_id=batch_id,
            status=AdmissionStatus.CONFIRMED,
            initial_payment_amount=15000.0,
            agreed_fee=25000.0,
            photo_bytes=b"dummy_bytes",
            admin_pin="1234",
            created_by=1,
        )
        adm_id = adm_service.create_admission(adm_dto)
        adm = adm_service.get_admission(adm_id)
        student_ids.append(adm.student_id)

    try:
        yield {
            "course_id": course_id,
            "batch_id": batch_id,
            "student_ids": student_ids,
            "names": names,
        }
    finally:
        repo = course_service.course_repo
        try:
            repo.execute("DELETE FROM attendance_records WHERE batch_id = ?", (batch_id,))
            repo.execute("DELETE FROM attendance_sessions WHERE batch_id = ?", (batch_id,))
            for sid in student_ids:
                repo.execute("DELETE FROM payments WHERE student_id = ?", (sid,))
                repo.execute("DELETE FROM receipts WHERE student_id = ?", (sid,))
                repo.execute("DELETE FROM admission_courses WHERE admission_id IN (SELECT id FROM admissions WHERE student_id = ?)", (sid,))
                repo.execute("DELETE FROM admissions WHERE student_id = ?", (sid,))
                repo.execute("DELETE FROM students WHERE id = ?", (sid,))
            repo.execute("DELETE FROM batches WHERE id = ?", (batch_id,))
            repo.execute("DELETE FROM courses WHERE id = ?", (course_id,))
            repo.execute("DELETE FROM activity_logs WHERE (entity_type = 'COURSE' AND entity_id = ?) OR (entity_type = 'BATCH' AND entity_id = ?)", (course_id, batch_id))
        except Exception:
            pass


def test_batch_enrolled_students_retrieval(setup_attendance_fixture):
    """
    Verifies that get_enrolled_students correctly retrieves all enrolled students
    for a batch with accurate candidate number, admission date, and pending fees.
    """
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]

    students = ctrl.get_enrolled_students(batch_id)
    assert len(students) == 3
    assert all(isinstance(s, BatchStudentDTO) for s in students)

    names = [s.student_name for s in students]
    assert "Aarav Mehta" in names
    assert "Neha Patil" in names
    assert "Vikram Rathore" in names

    first = students[0]
    assert first.student_id in setup_attendance_fixture["student_ids"]
    assert first.candidate_number != ""
    assert first.pending_fees == 10000.0
    assert first.admission_date == date.today().isoformat()


def test_online_attendance_recording_and_idempotency(setup_attendance_fixture):
    """
    Verifies online attendance recording using toggles (Default OFF = Absent, ON = Present),
    calculates present/absent counts accurately, and updates idempotently on the same date.
    """
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]
    session_date = "2026-10-15"

    # Step 1: Aarav is Present, Neha is Absent, Vikram is Present
    student_marks_1 = [
        {"student_id": stu_ids[0], "is_present": True},
        {"student_id": stu_ids[1], "is_present": False},
        {"student_id": stu_ids[2], "is_present": True},
    ]

    session_1 = ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date=session_date,
        student_marks=student_marks_1,
        recorded_by="admin",
        remarks="Morning online session",
    )

    assert session_1 is not None
    assert session_1.mode == AttendanceMode.ONLINE.value
    assert session_1.session_date == session_date
    assert session_1.total_enrolled == 3
    assert session_1.total_present == 2
    assert session_1.total_absent == 1

    # Check record details
    status_map = {r.student_id: r.status for r in session_1.records}
    assert status_map[stu_ids[0]] == "PRESENT"
    assert status_map[stu_ids[1]] == "ABSENT"
    assert status_map[stu_ids[2]] == "PRESENT"

    # Step 2: Idempotent update: teacher edits session, marks Neha Present too
    student_marks_2 = [
        {"student_id": stu_ids[0], "is_present": True},
        {"student_id": stu_ids[1], "is_present": True},
        {"student_id": stu_ids[2], "is_present": True},
    ]

    session_2 = ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date=session_date,
        student_marks=student_marks_2,
        recorded_by="admin",
        remarks="Corrected Neha status",
    )

    assert session_2.id == session_1.id  # Same session updated, not duplicated
    assert session_2.total_present == 3
    assert session_2.total_absent == 0


def test_offline_format_generation_with_customizer_columns(setup_attendance_fixture, tmp_path):
    """
    Verifies generation of fixed Excel attendance format with administrator column toggles.
    """
    service = AttendanceService()
    batch_id = setup_attendance_fixture["batch_id"]

    # 1. Customizer config: include pending fees, exclude admission date and contact
    config = OfflineFormatConfigDTO(
        batch_id=batch_id,
        include_student_name=True,
        include_admission_date=False,
        include_pending_fees=True,
        include_batch_details=True,
        include_course_details=True,
        include_contact_number=False,
        session_date="2026-10-18",
    )

    excel_path = service.generate_offline_format_excel(config, target_dir=tmp_path)
    assert os.path.exists(excel_path)
    assert excel_path.suffix == ".xlsx"

    # Open with openpyxl to verify content and columns
    import openpyxl

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws.title == "Attendance Sheet"

    # Verify header row at index 5
    headers = [cell.value for cell in ws[5] if cell.value is not None]
    assert "Student Name" in headers
    assert "Pending Fees" in headers
    assert "Adm Date" not in headers        # Excluded by admin toggle
    assert "Mobile Number" not in headers    # Excluded by admin toggle

    # Verify student data rows count starting at index 6
    data_rows = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[0] is not None]
    assert len(data_rows) == 3


def test_ocr_sheet_scanning_mark_detection(tmp_path, setup_attendance_fixture):
    """
    Generates a synthetic marked physical attendance sheet image,
    runs AttendanceOCREngine, and verifies mark detection.
    """
    batch_id = setup_attendance_fixture["batch_id"]
    students = AttendanceService().get_enrolled_students(batch_id)

    # Standard sheet dimensions
    img_width, img_height = 800, 600
    image = Image.new("RGB", (img_width, img_height), color=(255, 255, 255))
    draw = ImageDraw.Draw(image)

    # Header section (0 - 120px)
    draw.text((30, 20), "SIMS PHYSICAL ATTENDANCE SHEET", fill=(0, 0, 0))
    draw.text((30, 45), f"Batch ID: {batch_id}", fill=(0, 0, 0))

    # Grid parameters matching OCR engine layout
    table_start_y = int(img_height * 0.20)  # 120
    row_height = 48
    box_x = int(img_width * 0.80)           # 640

    # Mark student 0 and 2 with an ink mark, leave student 1 blank
    for idx, stu in enumerate(students):
        y = table_start_y + (idx * row_height)
        # Student info text
        draw.text((50, y + 10), f"{stu.candidate_number} | {stu.student_name}", fill=(0, 0, 0))

        # Checkbox outline
        draw.rectangle([box_x, y + 8, box_x + 36, y + 40], outline=(0, 0, 0), width=2)

        if idx in (0, 2):
            # Fill checkbox with black ink representing a marked physical sheet
            draw.rectangle([box_x + 3, y + 11, box_x + 33, y + 37], fill=(0, 0, 0))

    test_image_path = tmp_path / "offline_scan_test.png"
    image.save(test_image_path, format="PNG")

    # Run OCR Engine
    scan_result: OCRScanResultDTO = AttendanceOCREngine.scan_offline_sheet(
        image_path=test_image_path,
        batch_id=batch_id,
        batch_name="Morning Batch Alpha",
        enrolled_students=students,
        session_date="2026-10-18",
    )

    assert isinstance(scan_result, OCRScanResultDTO)
    assert scan_result.total_detected == 3
    assert len(scan_result.items) == 3

    # Check detected marks
    marks = {r.student_id: r.recognized_status for r in scan_result.items}
    assert marks[students[0].student_id] == "PRESENT"
    assert marks[students[1].student_id] == "ABSENT"
    assert marks[students[2].student_id] == "PRESENT"


def test_offline_attendance_reconciliation_and_commit(setup_attendance_fixture, tmp_path):
    """
    Verifies that offline attendance can be reconciled and committed to the database,
    storing marking_method='OCR_SCAN' and linking to the uploaded image.
    """
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]
    session_date = "2026-10-16"

    # Simulated verified marks from OCR reconciliation modal
    student_marks = [
        {"student_id": stu_ids[0], "status": "PRESENT", "confidence": 0.95},
        {"student_id": stu_ids[1], "status": "PRESENT", "confidence": 0.90},
        {"student_id": stu_ids[2], "status": "ABSENT", "confidence": 0.99},
    ]

    dummy_image = tmp_path / "sheet_scan_01.jpg"
    dummy_image.write_bytes(b"DUMMY_IMAGE_DATA")

    session = ctrl.commit_offline_attendance(
        batch_id=batch_id,
        session_date=session_date,
        student_marks=student_marks,
        scanned_image_path=str(dummy_image),
        recorded_by="admin",
        remarks="Scanned physical format sheet via OCR",
    )

    assert session is not None
    assert session.mode == AttendanceMode.OFFLINE.value
    assert session.scanned_image_path == str(dummy_image)
    assert session.total_present == 2
    assert session.total_absent == 1

    for r in session.records:
        assert r.marking_method == MarkingMethod.OCR_SCAN.value
        assert r.confidence_score > 0.8


def test_unified_batch_spreadsheet_matrix(setup_attendance_fixture):
    """
    Verifies that both online and offline attendance data flow into a unified
    batch spreadsheet matrix with per-session status, overall present count,
    absent count, and attendance percentage.
    """
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]

    # 1. Record an online session
    ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date="2026-10-15",
        student_marks=[{"student_id": s, "is_present": True} for s in stu_ids],
        recorded_by="admin",
    )

    # 2. Record an offline session
    ctrl.commit_offline_attendance(
        batch_id=batch_id,
        session_date="2026-10-16",
        student_marks=[{"student_id": s, "status": "PRESENT" if idx != 1 else "ABSENT"} for idx, s in enumerate(stu_ids)],
        recorded_by="admin",
    )

    # 3. Retrieve unified matrix
    matrix = ctrl.get_batch_matrix(batch_id)
    assert isinstance(matrix, BatchAttendanceMatrixDTO)
    assert matrix.batch_id == batch_id
    assert len(matrix.rows) == 3

    assert "2026-10-15" in matrix.session_dates
    assert "2026-10-16" in matrix.session_dates

    for row in matrix.rows:
        assert "2026-10-15" in row.date_marks
        assert "2026-10-16" in row.date_marks
        assert row.total_sessions >= 2
        assert row.attendance_percentage >= 0.0


def test_spreadsheet_excel_and_watermarked_pdf_exports(setup_attendance_fixture, tmp_path):
    """
    Verifies export of batch attendance spreadsheet to Excel (.xlsx)
    and Cairo vector PDF with diagonal translucent 'SIMS' watermark.
    """
    ctrl = AttendanceController()
    service = AttendanceService()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]

    # Record at least 1 session
    ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date="2026-10-17",
        student_marks=[{"student_id": s, "is_present": True} for s in stu_ids],
        recorded_by="admin",
    )

    # 1. Export Excel Matrix
    excel_path = service.export_matrix_excel(batch_id, target_dir=tmp_path)
    assert os.path.exists(excel_path)
    assert excel_path.suffix == ".xlsx"

    import openpyxl
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert "Attendance Matrix" in ws.title or "Sheet" in ws.title

    # 2. Export Cairo Vector PDF with Watermark
    pdf_path = service.export_matrix_pdf(batch_id, target_dir=tmp_path)
    assert os.path.exists(pdf_path)
    assert pdf_path.suffix == ".pdf"
    assert pdf_path.stat().st_size > 1000  # Non-trivial PDF generated

    # Verify PDF content contains PDF header markers
    pdf_bytes = pdf_path.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-")


def test_student_attendance_health_indicators(setup_attendance_fixture):
    """
    Verifies student attendance health calculations:
    Percentage, Total Sessions, Present/Absent counts, Health Status (EXCELLENT, GOOD, WARNING, CRITICAL),
    and session history timeline.
    """
    ctrl = AttendanceController()
    service = AttendanceService()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]

    # Record 2 sessions where student 0 is 100% present
    ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date="2026-10-15",
        student_marks=[{"student_id": stu_ids[0], "is_present": True}],
        recorded_by="admin",
    )
    ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date="2026-10-16",
        student_marks=[{"student_id": stu_ids[0], "is_present": True}],
        recorded_by="admin",
    )

    # Student 0 was present in both sessions -> 100% attendance -> EXCELLENT
    health_0 = service.get_student_attendance_health(stu_ids[0])
    assert isinstance(health_0, StudentAttendanceHealthDTO)
    assert health_0.total_sessions_conducted >= 2
    assert health_0.total_present >= 2
    assert health_0.total_absent == 0
    assert health_0.attendance_percentage == 100.0
    assert health_0.health_status == AttendanceHealth.EXCELLENT.value

    # Timeline records exist
    assert len(health_0.recent_history) >= 2
    for record in health_0.recent_history:
        assert record["status"] == "PRESENT"


def test_navigation_and_ui_view_mounting():
    """
    Verifies navigation rail destination, router registration for /attendance,
    and ensures DataTables strictly maintain non-empty columns in all states.
    """
    nav = NavigationMenu(on_nav_change=lambda _: None)
    assert "/attendance" in nav.routes

    # Check router protection
    assert "/attendance" in AppRouter.PROTECTED_ROUTES

    # Instantiate AttendanceHome
    view = AttendanceHome()
    assert view is not None
    assert hasattr(view, "load_data")
    assert hasattr(view, "load_batch_attendance")
    assert hasattr(view, "handle_save_online_attendance")
    assert hasattr(view, "handle_open_print_format")
    assert hasattr(view, "handle_open_ocr_upload")
    assert hasattr(view, "handle_export_excel")
    assert hasattr(view, "handle_export_pdf")

    # Critical Flutter DataTable Invariant: columns must contain at minimum one visible DataColumn
    assert len(view.matrix_table.columns) >= 1
    assert len(view.online_table.columns) >= 1
    assert len(view.health_table.columns) >= 1

    # Verify load_data keeps columns non-empty
    view.load_data()
    assert len(view.matrix_table.columns) >= 1
    assert len(view.online_table.columns) >= 1
    assert len(view.health_table.columns) >= 1

    # Verify empty matrix keeps columns non-empty
    empty_matrix = BatchAttendanceMatrixDTO(
        batch_id=999,
        batch_name="Empty Batch",
        batch_timing="—",
        course_name="—",
        session_dates=[],
        rows=[],
    )
    view._build_matrix_table(empty_matrix)
    assert len(view.matrix_table.columns) >= 1
    view._build_health_table(empty_matrix)
    assert len(view.health_table.columns) >= 1


def test_attendance_home_kpis_and_date_steppers(setup_attendance_fixture):
    """
    Verifies top KPI metrics cards initialization and quick date stepper functionality
    (< previous day, > next day, Today jump).
    """
    view = AttendanceHome()
    view.load_data()

    # Initial KPI card default values
    assert view.kpi_session_status_text.value != ""
    assert "%" in view.kpi_overall_rate_text.value
    assert "Sessions" in view.kpi_total_sessions_text.value
    assert "Students" in view.kpi_defaulters_text.value or "Defaulters" in view.kpi_defaulters_text.value

    # Test Date Stepper
    view.date_field.value = "2026-10-15"
    view._handle_prev_day(None)
    assert view.date_field.value == "2026-10-14"

    view._handle_next_day(None)
    assert view.date_field.value == "2026-10-15"

    view._handle_jump_today(None)
    assert view.date_field.value == date.today().isoformat()

    # Test DatePicker handler presence and safe execution
    assert hasattr(view, "datepicker_btn")
    assert hasattr(view, "_handle_open_datepicker")
    view._handle_open_datepicker(None)


def test_attendance_home_online_search_filter_and_invert(setup_attendance_fixture):
    """
    Verifies Tab 0 Daily Attendance table search, status filter, and invert marks batch action.
    """
    view = AttendanceHome()
    batch_id = setup_attendance_fixture["batch_id"]

    view.batch_dropdown.value = str(batch_id)
    view.load_batch_attendance(batch_id)

    # 3 enrolled students
    assert len(view.enrolled_students) == 3
    assert len(view.online_table.rows) == 3

    # Test Invert Marks: initially all False -> invert -> all True
    initial_switches = [s.value for s in view.student_toggles.values()]
    assert all(v is False for v in initial_switches)

    view.handle_invert_marks(None)
    assert all(s.value is True for s in view.student_toggles.values())

    # Test Search filter (case-insensitive substring)
    view.online_search_box.value = "Aarav"
    view._handle_online_search_change(None)
    assert len(view.online_table.rows) == 1

    view.online_search_box.value = ""
    view._handle_online_search_change(None)
    assert len(view.online_table.rows) == 3

    # Test Status filter: all are currently Present (True)
    view.online_filter_dropdown.value = "PRESENT"
    view._handle_online_filter_change(None)
    assert len(view.online_table.rows) == 3

    view.online_filter_dropdown.value = "ABSENT"
    view._handle_online_filter_change(None)
    assert len(view.online_table.rows) == 0

    # Reset filter
    view.online_filter_dropdown.value = "ALL"
    view._handle_online_filter_change(None)
    assert len(view.online_table.rows) == 3


def test_attendance_home_defaulters_and_health_tab(setup_attendance_fixture):
    """
    Verifies Tab 2 Defaulters & Student Health analytics table, search, health filter,
    and critical KPI count calculation.
    """
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    stu_ids = setup_attendance_fixture["student_ids"]

    # Record 3 sessions where student 0 is 100% present, students 1 and 2 are 100% absent (<60% -> CRITICAL)
    for s_date in ["2026-10-10", "2026-10-11", "2026-10-12"]:
        ctrl.record_online_attendance(
            batch_id=batch_id,
            session_date=s_date,
            student_marks=[
                {"student_id": stu_ids[0], "is_present": True},
                {"student_id": stu_ids[1], "is_present": False},
                {"student_id": stu_ids[2], "is_present": False},
            ],
            recorded_by="admin",
        )

    view = AttendanceHome()
    view.batch_dropdown.value = str(batch_id)
    view.load_batch_attendance(batch_id)

    # Invariant: health_table columns >= 1
    assert len(view.health_table.columns) >= 1
    assert len(view.health_table.rows) == 3

    # Defaulter count KPI should show 2 critical defaulters (<60%)
    assert "2 Students" in view.kpi_defaulters_text.value

    # Health Filter: Filter by CRITICAL
    view.health_filter_dropdown.value = "CRITICAL"
    view._handle_health_filter_change(None)
    assert len(view.health_table.rows) == 2

    # Health Filter: Filter by EXCELLENT
    view.health_filter_dropdown.value = "EXCELLENT"
    view._handle_health_filter_change(None)
    assert len(view.health_table.rows) == 1

    # Search in Health Tab
    view.health_filter_dropdown.value = "ALL"
    view._handle_health_filter_change(None)
    assert len(view.health_table.rows) == 3

    view.health_search_box.value = "Neha"
    view._handle_health_search_change(None)
    assert len(view.health_table.rows) == 1


def test_attendance_home_two_stage_batch_selection_and_back_button(setup_attendance_fixture):
    """
    Verifies that selecting a batch hides the batch selection card and reveals the active
    workspace with smooth transition, and clicking 'Change Batch' returns to the selector.
    """
    view = AttendanceHome()
    batch_id = setup_attendance_fixture["batch_id"]

    # Initial state before loading any batch
    assert hasattr(view, "batch_selection_card")
    assert hasattr(view, "active_batch_workspace")
    assert hasattr(view, "back_to_batches_btn")

    # Load batch attendance -> should transition into active workspace
    view.batch_dropdown.value = str(batch_id)
    view.load_batch_attendance(batch_id)

    assert view.batch_selection_card.visible is False
    assert view.active_batch_workspace.visible is True
    assert view.selected_batch_id == batch_id
    assert "Morning Batch" in view.active_batch_title_text.value

    # Click Back button -> transitions back to batch selection card
    view._handle_back_to_batches(None)

    assert view.batch_selection_card.visible is True
    assert view.active_batch_workspace.visible is False
    assert view.selected_batch_id is None


def test_future_date_attendance_validation_and_ui_blocking(setup_attendance_fixture):
    """
    Validates that:
    1. Service layer rejects future dates with ValidationError.
    2. AttendanceHome UI next_day button cannot advance past today.
    3. AttendanceHome UI date input cannot mark attendance for future dates.
    4. Save button rejects future dates cleanly.
    """
    from datetime import timedelta
    from core.exceptions import ValidationError
    from modules.attendance.repository import AttendanceRepository
    from modules.attendance.views.excel_import_reconciliation_modal import ExcelImportReconciliationModal

    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    students = ctrl.get_enrolled_students(batch_id)

    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    student_marks = [{"student_id": s.student_id, "admission_id": s.admission_id, "is_present": True} for s in students]

    # 1. Service layer must reject future attendance
    with pytest.raises(ValidationError) as exc_info:
        ctrl.record_online_attendance(
            batch_id=batch_id,
            session_date=tomorrow,
            student_marks=student_marks,
        )
    assert "future dates" in str(exc_info.value).lower()

    # 2. View level next day button on today's date should be blocked
    view = AttendanceHome()
    view.selected_batch_id = batch_id
    view.selected_date = date.today().isoformat()
    view.date_field.value = view.selected_date

    view._handle_next_day(None)
    assert view.selected_date == date.today().isoformat()

    # 3. View level handle_date_change with future date should clamp to today
    view.date_field.value = tomorrow
    view.handle_date_change(None)
    assert view.selected_date == date.today().isoformat()
    assert view.date_field.value == date.today().isoformat()

    # 4. View level save on future date should be blocked without throwing
    view.selected_date = tomorrow
    view.handle_save_online_attendance()
    # Toast shown, save not executed


def test_offline_upload_and_repository_posixpath_support(setup_attendance_fixture, tmp_path):
    """
    Verifies that passing a pathlib.Path to repository upsert_session and service commit_offline_attendance
    binds cleanly without SQLite parameter binding errors, and ExcelImportReconciliationModal accepts both on_saved and on_success.
    """
    from modules.attendance.repository import AttendanceRepository
    from modules.attendance.views.excel_import_reconciliation_modal import ExcelImportReconciliationModal

    service = AttendanceService()
    repo = AttendanceRepository()
    ctrl = AttendanceController()
    batch_id = setup_attendance_fixture["batch_id"]
    students = ctrl.get_enrolled_students(batch_id)

    # Passing pathlib.Path object for scanned_image_path
    image_path = tmp_path / "offline_sheet.png"
    image_path.write_bytes(b"\x89PNG\r\n\x1a\n")

    marks = [{"student_id": s.student_id, "admission_id": s.admission_id, "is_present": True} for s in students]

    # Test commit_offline_attendance with PosixPath
    session = service.commit_offline_attendance(
        batch_id=batch_id,
        session_date="2026-10-14",
        student_marks=marks,
        scanned_image_path=image_path,  # PosixPath instance
    )
    assert session is not None
    assert str(image_path) in (session.scanned_image_path or "")

    # Also test direct repo upsert_session with PosixPath within unit_of_work
    with service.unit_of_work():
        session_id = repo.upsert_session(
            batch_id=batch_id,
            session_date="2026-10-14",
            mode="OFFLINE",
            recorded_by="admin",
            remarks="Direct repo test",
            scanned_image_path=image_path,
        )
        assert session_id is not None
        assert session_id > 0

    # Verify ExcelImportReconciliationModal initializes with both on_success and on_saved
    callback_called = False
    def my_callback():
        nonlocal callback_called
        callback_called = True

    modal_success = ExcelImportReconciliationModal(
        batch_id=batch_id,
        batch_name="Test Batch",
        on_success=my_callback,
    )
    assert modal_success.on_saved is not None
    modal_success.on_saved()
    assert callback_called is True



