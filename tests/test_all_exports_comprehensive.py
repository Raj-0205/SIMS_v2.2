# tests/test_all_exports_comprehensive.py

from __future__ import annotations
import csv
from datetime import datetime
from pathlib import Path
import pytest
import openpyxl

from modules.reports.service import ReportsService
from modules.student.service import StudentService
from modules.student.dto import StudentFilterDTO
from modules.admission.service import AdmissionService
from modules.admission.dto import AdmissionFilterDTO
from modules.enquiry.service import EnquiryService
from modules.enquiry.dto import EnquiryFilterDTO


def test_reports_fees_exports(tmp_path: Path):
    service = ReportsService()
    
    excel_path = tmp_path / "fees_report.xlsx"
    csv_path = tmp_path / "fees_report.csv"
    
    res_excel = service.export_filtered_fees_report_excel(target_path=excel_path)
    assert Path(res_excel).exists()
    assert Path(res_excel).suffix == ".xlsx"
    
    # Verify openpyxl inspection
    wb = openpyxl.load_workbook(res_excel)
    assert "Fee Report" in wb.sheetnames
    ws = wb["Fee Report"]
    
    # Verify Title Banner
    assert "SUDHARM INFOTECH" in str(ws.cell(row=1, column=1).value)
    
    # Header row is at row 4
    header_row_idx = 4
    headers = [ws.cell(row=header_row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    assert "SR.NO" in headers
    assert "COURSE NAME" in headers
    assert "NAME" in headers
    assert "MOB NO" in headers
    assert "1ST INSTALLMENT" in headers
    assert "TOTAL FEES PAID" in headers
    assert "PENDING FEES" in headers
    
    # Check header styling (Navy blue fill, bold white font)
    header_cell = ws.cell(row=header_row_idx, column=1)
    assert header_cell.font.bold is True
    
    # Verify CSV export
    res_csv = service.export_filtered_fees_report_csv(target_path=csv_path)
    assert Path(res_csv).exists()
    
    with open(res_csv, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        first_row = next(reader)
        # Verify row 1 is headers without # metadata displacement
        assert first_row[0] == "SR.NO"
        assert "NAME" in first_row
        assert "MOB NO" in first_row
        assert "1ST INSTALLMENT" in first_row


def test_student_directory_exports(tmp_path: Path):
    service = StudentService()
    
    excel_path = tmp_path / "students.xlsx"
    csv_path = tmp_path / "students.csv"
    
    res_excel = service.export_students_excel(StudentFilterDTO(), target_path=excel_path)
    assert Path(res_excel).exists()
    
    wb = openpyxl.load_workbook(res_excel)
    ws = wb["Students"]
    assert "STUDENT DIRECTORY" in str(ws.cell(row=1, column=1).value)
    
    # Header row is at row 4
    header_row_idx = 4
    headers = [ws.cell(row=header_row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Student Name" in headers
    assert "Mobile Number" in headers
    assert "Current Course" in headers
    assert "Base Fee" in headers
    
    res_csv = service.export_students_csv(StudentFilterDTO(), target_path=csv_path)
    assert Path(res_csv).exists()
    with open(res_csv, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        first_row = next(reader)
        assert first_row[0] == "ID"
        assert "Student Name" in first_row
        assert "Mobile Number" in first_row


def test_admission_directory_exports(tmp_path: Path):
    service = AdmissionService()
    
    excel_path = tmp_path / "admissions.xlsx"
    csv_path = tmp_path / "admissions.csv"
    
    res_excel = service.export_admissions_excel(AdmissionFilterDTO(), target_path=excel_path)
    assert Path(res_excel).exists()
    
    wb = openpyxl.load_workbook(res_excel)
    ws = wb["Admissions"]
    assert "ADMISSIONS DIRECTORY" in str(ws.cell(row=1, column=1).value)
    
    # Header row is at row 4
    header_row_idx = 4
    headers = [ws.cell(row=header_row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    assert "SR.NO" in headers
    assert "COURSE NAME" in headers
    assert "NAME" in headers
    assert "MOB NO" in headers
    assert "1ST INSTALLMENT" in headers
    assert "TOTAL FEES PAID" in headers
    
    res_csv = service.export_admissions_csv(AdmissionFilterDTO(), target_path=csv_path)
    assert Path(res_csv).exists()
    with open(res_csv, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        first_row = next(reader)
        assert first_row[0] == "SR.NO"
        assert "NAME" in first_row
        assert "MOB NO" in first_row


def test_enquiry_exports(tmp_path: Path):
    service = EnquiryService()
    
    excel_path = tmp_path / "enquiries.xlsx"
    csv_path = tmp_path / "enquiries.csv"
    
    res_excel = service.export_enquiries_excel(EnquiryFilterDTO(), target_path=excel_path)
    assert Path(res_excel).exists()
    
    wb = openpyxl.load_workbook(res_excel)
    ws = wb["Enquiries"]
    assert "PROSPECTIVE STUDENT ENQUIRIES" in str(ws.cell(row=1, column=1).value)
    
    # Header row is at row 4
    header_row_idx = 4
    headers = [ws.cell(row=header_row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    assert "SR.NO" in headers
    assert "STUDENT NAME" in headers
    assert "MOBILE NO" in headers
    assert "INTERESTED COURSE" in headers
    assert "STATUS" in headers
    
    res_csv = service.export_enquiries_csv(EnquiryFilterDTO(), target_path=csv_path)
    assert Path(res_csv).exists()
    with open(res_csv, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        first_row = next(reader)
        assert first_row[0] == "SR.NO"
        assert "STUDENT NAME" in first_row
        assert "MOBILE NO" in first_row
