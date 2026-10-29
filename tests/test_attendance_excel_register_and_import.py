# tests/test_attendance_excel_register_and_import.py

from __future__ import annotations
import os
from datetime import date, datetime, timedelta
from pathlib import Path
import pytest
import openpyxl

from core.database.engine import DatabaseEngine
from modules.attendance.constants import (
    AttendanceMode,
    AttendanceStatus,
)
from modules.attendance.dto import (
    OfflineRegisterConfigDTO,
    ExcelAttendanceImportPreviewDTO,
    ExcelAttendanceCellDiffDTO,
    BatchStudentDTO,
)
from modules.attendance.service import AttendanceService
from modules.attendance.controller import AttendanceController
from modules.attendance.views.attendance_home import AttendanceHome
from modules.attendance.views.offline_register_dialog import OfflineRegisterDialog
from modules.attendance.views.excel_import_reconciliation_modal import ExcelImportReconciliationModal
from shared.utils.formatting import format_whatsapp_url


@pytest.fixture
def setup_excel_attendance_fixture():
    """
    Sets up an isolated course, batch, and 3 enrolled students with candidate numbers and phone numbers.
    """
    from core.security.context import SecurityContext
    from core.security.roles import Role
    from modules.admission.service import AdmissionService
    from modules.admission.dto import AdmissionCreateDTO, AdmissionStatus
    from modules.course.service import CourseService
    from modules.course.dto import CourseCreateDTO
    from modules.batch.controller import BatchController

    SecurityContext.set_current_user(user_id=1, username="admin", role=Role.ADMINISTRATOR.value)

    course_service = CourseService()
    batch_ctrl = BatchController()
    adm_service = AdmissionService()

    code = f"EX{int(datetime.now().timestamp() * 1000) % 100000}"
    course_id = course_service.create_course(
        CourseCreateDTO(name="Python Data Science", code=code, base_fee=20000.0, duration="3 Months")
    )

    batch_id = batch_ctrl.create_batch({
        "course_id": course_id,
        "batch_name": f"FastTrack Batch {code}",
        "timing": "11:00 AM - 01:00 PM",
        "max_capacity": 25,
        "status": "OPEN",
    })

    student_ids = []
    names = ["Rohan Deshmukh", "Pooja Kulkarni", "Aditya Shinde"]
    for idx, name in enumerate(names, start=1):
        fn, ln = name.split()
        adm_dto = AdmissionCreateDTO(
            first_name=fn,
            last_name=ln,
            mobile_number=f"982300{idx:04d}",
            secondary_mobile=f"982400{idx:04d}",
            village="Pune",
            course_id=course_id,
            batch_id=batch_id,
            status=AdmissionStatus.CONFIRMED,
            initial_payment_amount=10000.0,
            agreed_fee=20000.0,
            photo_bytes=b"sample_photo_bytes",
            admin_pin="1234",
            created_by=1,
        )
        adm_id = adm_service.create_admission(adm_dto)
        adm = adm_service.get_admission(adm_id)
        student_ids.append(adm.student_id)

    return {
        "course_id": course_id,
        "batch_id": batch_id,
        "student_ids": student_ids,
        "names": names,
    }


def test_offline_register_generation_all_durations(setup_excel_attendance_fixture, tmp_path):
    """
    Verifies generation of digital registers for 8 days, 15 days, 30 days, and 60 days.
    Checks Excel metadata, date columns count, student rows, and formula structures.
    """
    ctrl = AttendanceController()
    batch_id = setup_excel_attendance_fixture["batch_id"]

    durations = [8, 15, 30, 60]
    duration_keys = {8: "8_DAYS", 15: "15_DAYS", 30: "1_MONTH", 60: "2_MONTHS"}
    start_date = "2026-11-01"

    for duration in durations:
        config = OfflineRegisterConfigDTO(
            batch_id=batch_id,
            duration_type=duration_keys[duration],
            num_days=duration,
            start_date=start_date,
            include_pending_fees=True,
            include_contact_number=True,
        )

        out_path = ctrl.generate_offline_register(config, target_dir=tmp_path)
        assert Path(out_path).exists()
        assert Path(out_path).suffix == ".xlsx"

        # Load openpyxl workbook to inspect contents
        wb = openpyxl.load_workbook(out_path, data_only=False)
        ws = wb.active
        assert ws is not None

        # Verify header titles
        assert "ATTENDANCE REGISTER" in str(ws["A1"].value).upper()

        # Check date headers
        # Row 5 contains headers: A5=S.No, B5=Candidate No, C5=Student Name, D5=Mobile, E5=Pending Fees
        # Followed by date columns, then Present Days, Total Marked, Attendance %
        header_vals = [ws.cell(row=5, column=col).value for col in range(1, ws.max_column + 1)]
        assert "S.No" in header_vals
        assert "Student Name" in header_vals
        assert any("Present" in str(h) for h in header_vals)
        assert any("Total" in str(h) for h in header_vals)
        assert any("Attendance" in str(h) or "%" in str(h) for h in header_vals)

        # Check date headers count
        # Date columns start at col 6 (F)
        date_headers = []
        for col in range(6, 6 + duration):
            val = ws.cell(row=5, column=col).value
            if val and not any(k in str(val) for k in ("Present", "Total", "%")):
                date_headers.append(val)

        assert len(date_headers) == duration

        # Verify formula generation for first student (row 6)
        # Summary columns follow the duration columns
        p_col = 6 + duration
        t_col = 6 + duration + 1
        pct_col = 6 + duration + 2

        p_formula = str(ws.cell(row=6, column=p_col).value)
        t_formula = str(ws.cell(row=6, column=t_col).value)
        pct_formula = str(ws.cell(row=6, column=pct_col).value)

        assert "COUNTIF(" in p_formula
        assert '"P"' in p_formula
        assert "COUNTIF(" in t_formula
        assert "IF(" in pct_formula

        # Verify 3 students exist
        student_names = {ws.cell(row=r, column=3).value for r in (6, 7, 8)}
        assert student_names == {"Rohan Deshmukh", "Pooja Kulkarni", "Aditya Shinde"}

        wb.close()



def test_offline_register_parsing_and_non_destructive_blanks(setup_excel_attendance_fixture, tmp_path):
    """
    Verifies that uploaded registers correctly detect:
    1. Valid 'P' and 'A' markings.
    2. Non-destructive import: Blank/unmarked cells are strictly SKIPPED (never marked absent).
    3. Invalid markings are identified with warning.
    """
    ctrl = AttendanceController()
    batch_id = setup_excel_attendance_fixture["batch_id"]

    config = OfflineRegisterConfigDTO(
        batch_id=batch_id,
        duration_type="8_DAYS",
        num_days=8,
        start_date="2026-11-01",
    )
    reg_path = ctrl.generate_offline_register(config, target_dir=tmp_path)

    # Edit workbook to simulate teacher markings
    wb = openpyxl.load_workbook(reg_path)
    ws = wb.active

    # Student 1 (row 6): Day 1 = P, Day 2 = P, Day 3 = A, Days 4-8 = BLANK
    ws.cell(row=6, column=6).value = "P"
    ws.cell(row=6, column=7).value = "P"
    ws.cell(row=6, column=8).value = "A"
    for col in range(9, 14):
        ws.cell(row=6, column=col).value = ""  # Explicit blank

    # Student 2 (row 7): Day 1 = A, Day 2 = P, Day 3 = BLANK, Day 4 = INVALID "L"
    ws.cell(row=7, column=6).value = "A"
    ws.cell(row=7, column=7).value = "P"
    ws.cell(row=7, column=8).value = None  # None
    ws.cell(row=7, column=9).value = "L"   # Invalid status mark

    # Student 3 (row 8): All days BLANK (absentee or future dates)
    for col in range(6, 14):
        ws.cell(row=8, column=col).value = ""

    marked_path = tmp_path / "marked_register.xlsx"
    wb.save(marked_path)
    wb.close()

    # Parse and generate preview
    preview = ctrl.parse_offline_register(marked_path, batch_id=batch_id)

    assert isinstance(preview, ExcelAttendanceImportPreviewDTO)
    assert preview.batch_id == batch_id
    assert len(preview.detected_dates) == 8

    # Check action counts
    # Student 1: 3 new (P, P, A), 5 skipped
    # Student 2: 2 new (A, P), 1 invalid (L), 5 skipped
    # Student 3: 8 skipped
    # Total evaluated = 3 * 8 = 24
    assert preview.new_marks_count == 5
    assert preview.invalid_cells_count == 1
    assert preview.skipped_blank_count == 18
    assert preview.overwrite_marks_count == 0
    assert preview.requires_admin_pin is False

    # Check diff cell details
    invalid_cells = [c for c in preview.diff_items if c.change_type == "INVALID"]
    assert len(invalid_cells) == 1
    assert invalid_cells[0].student_name == "Pooja Kulkarni"
    assert invalid_cells[0].mark == "L"


def test_offline_register_commit_and_admin_pin_security(setup_excel_attendance_fixture, tmp_path):
    """
    Verifies transactional commit, percentage recalculation, and Admin PIN step-up enforcement
    when overwriting existing attendance records.
    """
    ctrl = AttendanceController()
    batch_id = setup_excel_attendance_fixture["batch_id"]
    stu_ids = setup_excel_attendance_fixture["student_ids"]

    # 1. First record an initial session for 2026-10-10 where all 3 students are ABSENT
    initial_date = "2026-10-10"
    ctrl.record_online_attendance(
        batch_id=batch_id,
        session_date=initial_date,
        student_marks=[
            {"student_id": stu_ids[0], "is_present": False},
            {"student_id": stu_ids[1], "is_present": False},
            {"student_id": stu_ids[2], "is_present": False},
        ],
    )

    # Verify initial matrix has 0% attendance
    matrix = ctrl.get_batch_matrix(batch_id)
    assert matrix.overall_attendance_rate == 0.0

    # 2. Generate register starting at 2026-10-10 for 8 days
    config = OfflineRegisterConfigDTO(
        batch_id=batch_id,
        duration_type="8_DAYS",
        num_days=8,
        start_date=initial_date,
    )
    reg_path = ctrl.generate_offline_register(config, target_dir=tmp_path)

    # 3. Mark Day 1 (2026-10-10) as PRESENT for Student 0 and Student 1 (OVERWRITE!)
    # Mark Day 1 as ABSENT for Student 2 (UNCHANGED)
    # And mark Day 2 (2026-10-11) as PRESENT for Student 0 and Student 1, ABSENT for Student 2 (NEW SESSION!)
    wb = openpyxl.load_workbook(reg_path)
    ws = wb.active

    # Dynamically find column indices for Day 1 (2026-10-10) and Day 2 (2026-10-11)
    col_day1 = None
    col_day2 = None
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(row=5, column=c).value or "")
        if "2026-10-10" in v:
            col_day1 = c
        elif "2026-10-11" in v:
            col_day2 = c

    assert col_day1 is not None, "Column for 2026-10-10 not found"
    assert col_day2 is not None, "Column for 2026-10-11 not found"

    # Dynamically map student names to rows
    row_by_name = {}
    for r in range(6, 9):
        name = ws.cell(row=r, column=3).value
        row_by_name[name] = r

    # Student 0 (Rohan Deshmukh): Day 1 A -> P (Overwrite), Day 2 -> P (New)
    r_rohan = row_by_name["Rohan Deshmukh"]
    ws.cell(row=r_rohan, column=col_day1).value = "P"
    ws.cell(row=r_rohan, column=col_day2).value = "P"

    # Student 1 (Pooja Kulkarni): Day 1 A -> P (Overwrite), Day 2 -> P (New)
    r_pooja = row_by_name["Pooja Kulkarni"]
    ws.cell(row=r_pooja, column=col_day1).value = "P"
    ws.cell(row=r_pooja, column=col_day2).value = "P"

    # Student 2 (Aditya Shinde): Day 1 A -> A (Unchanged), Day 2 -> A (New)
    r_aditya = row_by_name["Aditya Shinde"]
    ws.cell(row=r_aditya, column=col_day1).value = "A"
    ws.cell(row=r_aditya, column=col_day2).value = "A"

    overwrite_path = tmp_path / "overwrite_register.xlsx"
    wb.save(overwrite_path)
    wb.close()

    # Parse and inspect diff
    preview = ctrl.parse_offline_register(overwrite_path, batch_id=batch_id)
    assert preview.overwrite_marks_count == 2  # Student 0 and Student 1 changed from A to P
    assert preview.unchanged_marks_count == 1   # Student 2 was A and remains A
    assert preview.new_marks_count == 3         # 2026-10-11 has 3 marks
    assert preview.requires_admin_pin is True

    # 4. Attempt commit without PIN or with wrong PIN -> must raise error
    with pytest.raises(Exception) as exc_info:
        ctrl.commit_offline_register_import(preview, admin_pin="9999")
    assert "PIN" in str(exc_info.value) or "denied" in str(exc_info.value).lower()

    # 5. Commit with correct Admin PIN ("1234")
    committed_cnt = ctrl.commit_offline_register_import(preview, admin_pin="1234")
    assert committed_cnt > 0


    # 6. Verify matrix has recalculated with new attendance
    updated_matrix = ctrl.get_batch_matrix(batch_id)
    assert len(updated_matrix.session_dates) == 2
    assert "2026-10-10" in updated_matrix.session_dates
    assert "2026-10-11" in updated_matrix.session_dates

    # Check Student 0: 2 sessions, both Present -> 100% attendance
    r0 = next(r for r in updated_matrix.rows if r.student_id == stu_ids[0])
    assert r0.total_present == 2
    assert r0.attendance_percentage == 100.0

    # Overall batch attendance: 4 present out of 6 marks = 66.7%
    assert updated_matrix.overall_attendance_rate == 66.7


def test_offline_register_dialog_and_import_modal_ui(setup_excel_attendance_fixture):
    """
    Verifies frontend dialog and modal instantiations and invariant protection.
    """
    batch_id = setup_excel_attendance_fixture["batch_id"]

    # Test OfflineRegisterDialog
    reg_dlg = OfflineRegisterDialog(
        batch_id=batch_id,
        batch_name="Morning Python",
        course_name="Full Stack Python",
    )
    assert reg_dlg is not None
    assert reg_dlg.duration_dropdown.value == "15_DAYS"
    assert len(reg_dlg.duration_dropdown.options) == 4
    assert hasattr(reg_dlg, "handle_generate")

    # Test ExcelImportReconciliationModal
    import_modal = ExcelImportReconciliationModal(
        batch_id=batch_id,
        batch_name="Morning Python",
    )
    assert import_modal is not None
    assert hasattr(import_modal, "handle_browse_file")
    assert import_modal.diff_list is not None
    assert import_modal.admin_pin_input is not None


def test_attendance_home_subtabs_and_whatsapp_integration(setup_excel_attendance_fixture):
    """
    Verifies AttendanceHome 4-subtab layout, WhatsApp URL generation, and Admin PIN save flow.
    """
    batch_id = setup_excel_attendance_fixture["batch_id"]

    view = AttendanceHome(initial_batch_id=batch_id)
    view.load_data()

    # 4 Subtabs verified
    assert hasattr(view, "tab_online_btn")
    assert hasattr(view, "tab_offline_btn")
    assert hasattr(view, "tab_matrix_btn")
    assert hasattr(view, "tab_health_btn")

    # Invariant: all tables have columns >= 1
    assert len(view.online_table.columns) >= 1
    assert len(view.matrix_table.columns) >= 1
    assert len(view.health_table.columns) >= 1

    # Switch across all 4 subtabs
    view._set_subtab(0)
    assert view.online_panel.visible is True
    assert view.offline_panel.visible is False

    view._set_subtab(1)
    assert view.online_panel.visible is False
    assert view.offline_panel.visible is True

    view._set_subtab(2)
    assert view.matrix_panel.visible is True
    assert view.offline_panel.visible is False

    view._set_subtab(3)
    assert view.health_panel.visible is True
    assert view.matrix_panel.visible is False

    # Test WhatsApp URL formatting helper
    wa_url = format_whatsapp_url("9823000001", "Hello Rohan, you are marked ABSENT.")
    assert "https://wa.me/919823000001" in wa_url
    assert "ABSENT" in wa_url

    # Test Admin PIN save flow
    view.selected_batch_id = batch_id
    view.load_batch_attendance(batch_id)
    # Mark all present and save with admin PIN
    view.handle_mark_all_present(None)
    view.handle_save_online_attendance(None, admin_pin="1234")

    # Confirm session was recorded
    session = view.controller.get_session(batch_id, view.selected_date)
    assert session is not None
    assert session.total_present == 3
